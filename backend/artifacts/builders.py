"""Deterministic, host-owned builders for common deliverable formats."""

from __future__ import annotations

from dataclasses import dataclass
import base64
import csv
from datetime import datetime, timezone
from functools import lru_cache
from io import BytesIO, StringIO
import hashlib
import html
import json
import math
import os
import re
import zipfile
from typing import Any, Mapping, Sequence

from core_invariants import canonical_json_bytes, strict_json_value

BUILDER_VERSION = "variant1-artifact-builders.2"
_MAX_BLOCKS = 2_000
_MAX_ROWS = 50_000
_MAX_CELLS = 500_000
_MAX_TEXT = 2_000_000
_MAX_IMAGE_BYTES = 12 * 1024 * 1024
_BLOCK_TYPES = frozenset({
    "paragraph", "heading", "bullets", "table", "code", "quote", "page_break", "image",
})


@dataclass(frozen=True, slots=True)
class BuildOutput:
    format: str
    payload: bytes
    media_type: str
    extension: str
    renderer: str
    renderer_version: str
    normalized_sha256: str
    previews: tuple[bytes, ...] = ()
    diagnostics: tuple[Mapping[str, Any], ...] = ()


def _strict(value: Any) -> Any:
    clean = strict_json_value(value)
    encoded = canonical_json_bytes(clean)
    if len(encoded) > 16 * 1024 * 1024:
        raise ValueError("artifact specification exceeds 16 MiB")
    return clean


def freeze_spec_resources(spec: Mapping[str, Any]) -> dict[str, Any]:
    """Resolve file conveniences once, before the canonical spec enters CAS.

    Pure builders only receive embedded resources. This keeps path-based
    authoring useful while later renders never re-read changing source files.
    """
    clean = _strict(dict(spec))
    def visit(value):
        if isinstance(value, dict):
            if value.get('type') == 'image' and value.get('path'):
                if value.get('data_base64'):
                    raise ValueError('Image blocks must supply either path or data_base64, not both.')
                with open(os.path.expanduser(str(value.pop('path'))), 'rb') as stream:
                    payload = stream.read(_MAX_IMAGE_BYTES + 1)
                if len(payload) > _MAX_IMAGE_BYTES:
                    raise ValueError('Image exceeds the 12 MiB image bound.')
                value['data_base64'] = base64.b64encode(payload).decode('ascii')
                _, mime = _image_bytes(value)
                value['mime_type'] = mime
            for child in value.values():
                visit(child)
        elif isinstance(value, list):
            for child in value:
                visit(child)
    visit(clean)
    return _strict(clean)


def _text(value: Any) -> str:
    result = str(value or "")
    if len(result) > _MAX_TEXT:
        raise ValueError("artifact text exceeds bound")
    return result


def _blocks(spec: Mapping[str, Any]) -> list[dict[str, Any]]:
    rows = spec.get("blocks")
    if rows is None:
        rows = []
    if not isinstance(rows, list) or len(rows) > _MAX_BLOCKS:
        raise ValueError("blocks must be a bounded list")
    if any(not isinstance(row, Mapping) for row in rows):
        raise ValueError("every artifact block must be an object")
    return [dict(row) for row in rows]


def _image_bytes(block: Mapping[str, Any]) -> tuple[bytes, str]:
    """Decode one image from the saved JSON spec, never a mutable host path."""
    if block.get("path"):
        raise ValueError(
            "image.path is not reproducible; embed PNG/JPEG bytes as image.data_base64"
        )
    encoded = block.get("data_base64")
    if not isinstance(encoded, str) or not encoded:
        raise ValueError("image block requires non-empty data_base64 PNG/JPEG bytes")
    try:
        payload = base64.b64decode(encoded, validate=True)
    except Exception as exc:
        raise ValueError("image.data_base64 is not valid base64") from exc
    if not payload or len(payload) > _MAX_IMAGE_BYTES:
        raise ValueError("image bytes are empty or exceed the 12 MiB image bound")
    from PIL import Image

    try:
        with Image.open(BytesIO(payload)) as image:
            actual = str(image.format or "").upper()
            image.verify()
    except Exception as exc:
        raise ValueError("image bytes cannot be opened") from exc
    mime = {"PNG": "image/png", "JPEG": "image/jpeg"}.get(actual)
    if mime is None:
        raise ValueError(f"unsupported image format {actual or 'unknown'}; use PNG or JPEG")
    requested = str(block.get("mime_type") or mime).lower()
    if requested != mime:
        raise ValueError(f"image mime_type {requested} disagrees with its {actual} bytes")
    return payload, mime


def _image_width_in(block: Mapping[str, Any], default: float) -> float:
    try:
        width = float(block.get("width_in") or default)
    except (TypeError, ValueError) as exc:
        raise ValueError("image.width_in must be a finite number") from exc
    if not math.isfinite(width) or not 0.1 <= width <= 20:
        raise ValueError("image.width_in must be between 0.1 and 20 inches")
    return width


def _validate_content(spec: Mapping[str, Any], format_name: str) -> None:
    if spec.get("slides") is not None and format_name != "pptx":
        raise ValueError("slides are supported only by the PPTX builder; no slide content was rendered")
    if spec.get("slides") is not None and not isinstance(spec["slides"], list):
        raise ValueError("slides must be a list")
    groups = [_blocks(spec)]
    if format_name == "pptx" and spec.get("slides") is not None and groups[0]:
        raise ValueError("PPTX spec cannot mix top-level blocks with slides; move blocks into a slide")
    if format_name == "pptx" and isinstance(spec.get("slides"), list):
        groups.extend(_blocks({"blocks": dict(slide).get("blocks")})
                      for slide in spec["slides"] if isinstance(slide, Mapping))
        if any(not isinstance(slide, Mapping) for slide in spec["slides"]):
            raise ValueError("every slide must be an object")
    for blocks in groups:
        for block in blocks:
            kind = str(block.get("type") or "paragraph")
            if kind not in _BLOCK_TYPES:
                raise ValueError(f"unsupported artifact block type: {kind}")
            if kind == "image":
                _image_bytes(block)
                if format_name not in {"md", "markdown", "html", "docx", "pptx", "pdf", "json"}:
                    raise ValueError(f"{format_name} cannot render image blocks; choose PDF, DOCX, PPTX, HTML, or Markdown")


def _markdown(spec: Mapping[str, Any]) -> str:
    lines: list[str] = []
    title = _text(spec.get("title")).strip()
    if title:
        lines.extend((f"# {title}", ""))
    for block in _blocks(spec):
        kind = str(block.get("type") or "paragraph")
        if kind == "heading":
            level = max(1, min(int(block.get("level") or 2), 6))
            lines.extend(("#" * level + " " + _text(block.get("text")).strip(), ""))
        elif kind == "bullets":
            for item in list(block.get("items") or ())[:10_000]:
                lines.append(f"- {_text(item)}")
            lines.append("")
        elif kind == "table":
            rows = list(block.get("rows") or ())[:_MAX_ROWS]
            rows = [list(row) for row in rows if isinstance(row, (list, tuple))]
            if rows:
                width = max(len(row) for row in rows)
                header = rows[0] + [""] * (width - len(rows[0]))
                lines.append("| " + " | ".join(_text(item) for item in header) + " |")
                lines.append("| " + " | ".join("---" for _ in range(width)) + " |")
                for row in rows[1:]:
                    padded = row + [""] * (width - len(row))
                    lines.append("| " + " | ".join(_text(item) for item in padded) + " |")
                lines.append("")
        elif kind == "code":
            language = re.sub(r"[^A-Za-z0-9_+-]", "", str(block.get("language") or ""))
            lines.extend((f"```{language}", _text(block.get("text")), "```", ""))
        elif kind == "quote":
            lines.extend((*("> " + line for line in _text(block.get("text")).splitlines()), ""))
        elif kind == "image":
            payload, mime = _image_bytes(block)
            alt = _text(block.get("alt")).replace("]", r"\]")
            lines.extend((
                f"![{alt}](data:{mime};base64,{base64.b64encode(payload).decode('ascii')})", "",
            ))
        elif kind == "page_break":
            lines.extend(("---", ""))
        else:
            lines.extend((_text(block.get("text")), ""))
    return "\n".join(lines).rstrip() + "\n"


def _html(spec: Mapping[str, Any]) -> str:
    body: list[str] = []
    title = _text(spec.get("title")).strip()
    if title:
        body.append(f"<h1>{html.escape(title)}</h1>")
    for block in _blocks(spec):
        kind = str(block.get("type") or "paragraph")
        if kind == "heading":
            level = max(1, min(int(block.get("level") or 2), 6))
            body.append(f"<h{level}>{html.escape(_text(block.get('text')))}</h{level}>")
        elif kind == "bullets":
            body.append("<ul>" + "".join(
                f"<li>{html.escape(_text(item))}</li>"
                for item in list(block.get("items") or ())[:10_000]
            ) + "</ul>")
        elif kind == "table":
            rows = [list(row) for row in list(block.get("rows") or ())[:_MAX_ROWS]
                    if isinstance(row, (list, tuple))]
            if rows:
                body.append("<table><thead><tr>" + "".join(
                    f"<th>{html.escape(_text(item))}</th>" for item in rows[0]
                ) + "</tr></thead><tbody>" + "".join(
                    "<tr>" + "".join(
                        f"<td>{html.escape(_text(item))}</td>" for item in row
                    ) + "</tr>" for row in rows[1:]
                ) + "</tbody></table>")
        elif kind == "code":
            body.append(f"<pre><code>{html.escape(_text(block.get('text')))}</code></pre>")
        elif kind == "quote":
            body.append(f"<blockquote>{html.escape(_text(block.get('text')))}</blockquote>")
        elif kind == "image":
            payload, mime = _image_bytes(block)
            alt = html.escape(_text(block.get("alt")), quote=True)
            source = f"data:{mime};base64,{base64.b64encode(payload).decode('ascii')}"
            body.append(f'<figure><img src="{source}" alt="{alt}" style="max-width:100%"></figure>')
        elif kind == "page_break":
            body.append('<div style="break-after:page"></div>')
        else:
            body.append(f"<p>{html.escape(_text(block.get('text')))}</p>")
    return (
        "<!doctype html><html><head><meta charset=\"utf-8\"><title>"
        + html.escape(title or "VARIANT-1 artifact")
        + "</title><style>body{font:16px/1.5 system-ui;max-width:960px;margin:40px auto;padding:0 24px}"
          "table{border-collapse:collapse;width:100%}th,td{border:1px solid #bbb;padding:6px;text-align:left}"
          "pre{white-space:pre-wrap;background:#f3f3f3;padding:12px}</style></head><body>"
        + "".join(body) + "</body></html>"
    )


def _docx(spec: Mapping[str, Any]) -> bytes:
    from docx import Document
    from docx.shared import Inches

    document = Document()
    document.core_properties.title = _text(spec.get("title"))
    fixed = datetime(2000, 1, 1, tzinfo=timezone.utc)
    document.core_properties.created = fixed
    document.core_properties.modified = fixed
    title = _text(spec.get("title")).strip()
    if title:
        document.add_heading(title, 0)
    for block in _blocks(spec):
        kind = str(block.get("type") or "paragraph")
        if kind == "heading":
            document.add_heading(
                _text(block.get("text")),
                level=max(1, min(int(block.get("level") or 2), 9)),
            )
        elif kind == "bullets":
            for item in list(block.get("items") or ())[:10_000]:
                document.add_paragraph(_text(item), style="List Bullet")
        elif kind == "table":
            rows = [list(row) for row in list(block.get("rows") or ())[:_MAX_ROWS]
                    if isinstance(row, (list, tuple))]
            if rows:
                width = max(len(row) for row in rows)
                table = document.add_table(rows=len(rows), cols=width)
                table.style = "Table Grid"
                for r_index, row in enumerate(rows):
                    for c_index, item in enumerate(row):
                        table.cell(r_index, c_index).text = _text(item)
        elif kind == "page_break":
            document.add_page_break()
        elif kind == "image":
            payload, _mime = _image_bytes(block)
            document.add_picture(BytesIO(payload), width=Inches(_image_width_in(block, 6)))
        else:
            document.add_paragraph(_text(block.get("text")))
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pptx(spec: Mapping[str, Any]) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from PIL import Image

    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    slides = spec.get("slides")
    if not isinstance(slides, list):
        slides = [{"title": spec.get("title") or "", "blocks": _blocks(spec)}]
    if len(slides) > 500:
        raise ValueError("slide count exceeds bound")
    for raw in slides:
        slide_spec = dict(raw) if isinstance(raw, Mapping) else {}
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        title = slide.shapes.title
        if title is not None:
            title.text = _text(slide_spec.get("title"))
        blocks = _blocks(slide_spec)
        image_blocks = [block for block in blocks if block.get("type") == "image"]
        textbox = slide.shapes.add_textbox(
            Inches(.8), Inches(1.45), Inches(6.2 if image_blocks else 11.7), Inches(5.4),
        )
        frame = textbox.text_frame
        frame.word_wrap = True
        lines: list[tuple[str, int]] = []
        for block in blocks:
            if block.get("type") == "bullets":
                lines.extend((_text(item), 0) for item in list(block.get("items") or ()))
            elif block.get("type") == "table":
                lines.extend((" | ".join(_text(item) for item in row), 0)
                             for row in list(block.get("rows") or ()))
            elif block.get("type") == "image":
                continue
            elif block.get("type") == "page_break":
                raise ValueError("page_break is not a PPTX slide block; add a separate slide")
            else:
                lines.append((_text(block.get("text")), 0))
        if not lines and slide_spec.get("body"):
            lines = [(_text(slide_spec.get("body")), 0)]
        for index, (line, level) in enumerate(lines):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.level = level
            paragraph.font.size = Pt(22)
        image_y = 1.6
        for block in image_blocks:
            payload, _mime = _image_bytes(block)
            with Image.open(BytesIO(payload)) as image:
                pixels_w, pixels_h = image.size
            width = _image_width_in(block, 5)
            height = width * pixels_h / pixels_w
            if block.get("width_in") is None and height > 5.5:
                width *= 5.5 / height
                height = 5.5
            try:
                x = float(block.get("x_in") if block.get("x_in") is not None else 7.4)
                y = float(block.get("y_in") if block.get("y_in") is not None else image_y)
            except (TypeError, ValueError) as exc:
                raise ValueError("image x_in/y_in must be finite slide coordinates") from exc
            if not all(math.isfinite(value) for value in (x, y)) or (
                x < 0 or y < 1.3 or x + width > 13.333333 or y + height > 7.5
            ):
                raise ValueError("PPTX image placement exceeds the slide; adjust x_in/y_in/width_in")
            picture = slide.shapes.add_picture(
                BytesIO(payload), Inches(x), Inches(y), width=Inches(width),
            )
            if block.get("alt"):
                picture.name = _text(block.get("alt"))[:255]
            if block.get("y_in") is None:
                image_y = y + height + .2
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def _xlsx(spec: Mapping[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = spec.get("sheets")
    if not isinstance(sheets, list):
        sheets = [{"name": "Sheet1", "rows": spec.get("rows") or []}]
    if not sheets or len(sheets) > 200:
        raise ValueError("workbook requires 1-200 sheets")
    cell_count = 0
    for index, raw in enumerate(sheets):
        sheet_spec = dict(raw) if isinstance(raw, Mapping) else {}
        name = re.sub(r"[\\/*?:\[\]]", "_", _text(sheet_spec.get("name") or f"Sheet{index + 1}"))[:31]
        sheet = workbook.create_sheet(name or f"Sheet{index + 1}")
        rows = list(sheet_spec.get("rows") or ())
        if len(rows) > _MAX_ROWS:
            raise ValueError("worksheet row count exceeds bound")
        for r_index, raw_row in enumerate(rows, 1):
            row = list(raw_row) if isinstance(raw_row, (list, tuple)) else [raw_row]
            cell_count += len(row)
            if cell_count > _MAX_CELLS:
                raise ValueError("workbook cell count exceeds bound")
            for c_index, value in enumerate(row, 1):
                sheet.cell(r_index, c_index, value=value)
        if rows:
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2C3E50")
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(
                60, max(10, max(len(str(cell.value or "")) for cell in column) + 2)
            )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _delimited(spec: Mapping[str, Any], delimiter: str) -> bytes:
    rows = list(spec.get("rows") or ())
    if len(rows) > _MAX_ROWS:
        raise ValueError("row count exceeds bound")
    output = StringIO(newline="")
    writer = csv.writer(output, delimiter=delimiter, lineterminator="\n")
    count = 0
    for row in rows:
        values = list(row) if isinstance(row, (list, tuple)) else [row]
        count += len(values)
        if count > _MAX_CELLS:
            raise ValueError("cell count exceeds bound")
        writer.writerow([_text(item) for item in values])
    return output.getvalue().encode("utf-8-sig")


@lru_cache(maxsize=1)
def _pdf_font_pool() -> tuple[tuple[str, str, dict[int, int], dict[int, int]], ...]:
    """Register available outline fonts and retain their actual glyph maps."""
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    package_fonts = os.path.join(os.path.dirname(reportlab.__file__), "fonts")
    candidates = (
        (r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\msyhbd.ttc"),
        (r"C:\Windows\Fonts\malgun.ttf", r"C:\Windows\Fonts\malgunbd.ttf"),
        (r"C:\Windows\Fonts\meiryo.ttc", r"C:\Windows\Fonts\meiryob.ttc"),
        (r"C:\Windows\Fonts\segoeui.ttf", r"C:\Windows\Fonts\segoeuib.ttf"),
        (r"C:\Windows\Fonts\seguisym.ttf", r"C:\Windows\Fonts\seguisym.ttf"),
        ("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc", "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"),
        ("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc", "/usr/share/fonts/truetype/noto/NotoSansCJK-Bold.ttc"),
        ("/usr/share/fonts/noto-cjk/NotoSansCJK-Regular.ttc", "/usr/share/fonts/noto-cjk/NotoSansCJK-Bold.ttc"),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        (os.path.join(package_fonts, "Vera.ttf"), os.path.join(package_fonts, "VeraBd.ttf")),
    )
    loaded: list[tuple[str, str, dict[int, int], dict[int, int]]] = []
    for regular_path, bold_path in candidates:
        if not os.path.isfile(regular_path):
            continue
        index = len(loaded)
        regular_name = f"Variant1Unicode{index}"
        bold_name = f"Variant1UnicodeBold{index}"
        try:
            regular = TTFont(regular_name, regular_path)
            bold = TTFont(bold_name, bold_path if os.path.isfile(bold_path) else regular_path)
            pdfmetrics.registerFont(regular)
            pdfmetrics.registerFont(bold)
            loaded.append((
                regular_name, bold_name,
                dict(regular.face.charToGlyph), dict(bold.face.charToGlyph),
            ))
        except Exception:
            continue
    if not loaded:
        raise ValueError("PDF needs an installed Unicode outline font")
    return tuple(loaded)


def _pdf_glyph_runs(text: str, *, bold: bool) -> list[tuple[str, str]]:
    runs: list[tuple[str, str]] = []
    for character in text:
        codepoint = ord(character)
        if codepoint > 0xFFFF:
            # reportlab 5's TrueType ToUnicode emitter encodes these as an
            # odd-length glyph hex string (for example U+1F600), producing a
            # visible-but-unextractable emoji. Do not publish that loss as a
            # successful PDF; a raster image block can preserve the pixels.
            raise ValueError(
                f"PDF text cannot preserve supplementary U+{codepoint:05X}; provide it as an image block"
            )
        selected = ""
        for regular, strong, regular_map, bold_map in _pdf_font_pool():
            if bold and bold_map.get(codepoint):
                selected = strong
                break
            if regular_map.get(codepoint):
                selected = regular
                break
        if not selected:
            raise ValueError(
                f"PDF font coverage cannot represent U+{codepoint:04X}; install a font with that glyph or provide it as an image"
            )
        if runs and runs[-1][0] == selected:
            name, value = runs[-1]
            runs[-1] = name, value + character
        else:
            runs.append((selected, character))
    return runs


def _pdf(spec: Mapping[str, Any]) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfgen import canvas
    from PIL import Image

    _pdf_font_pool()
    buffer = BytesIO()
    page_width, page_height = A4
    pdf = canvas.Canvas(buffer, pagesize=A4, invariant=1, pageCompression=1)
    pdf.setTitle(_text(spec.get("title")) or "VARIANT-1 artifact")
    margin = 54
    y = page_height - margin
    content_width = page_width - 2 * margin

    def new_page() -> None:
        nonlocal y
        pdf.showPage()
        y = page_height - margin

    def draw_line(value: str, *, size: float, bold: bool) -> None:
        nonlocal y
        if y - size * 1.45 < margin:
            new_page()
        x = margin
        for font, run in _pdf_glyph_runs(value, bold=bold):
            pdf.setFont(font, size)
            pdf.drawString(x, y, run)
            x += pdfmetrics.stringWidth(run, font, size)
        y -= size * 1.45

    def draw_text(value: Any, *, size: float = 10, bold: bool = False) -> None:
        nonlocal y
        for source_line in _text(value).split("\n"):
            if not source_line:
                y -= size * 0.8
                continue
            line = ""
            line_width = 0.0
            for character in source_line.replace("\t", "    "):
                font = _pdf_glyph_runs(character, bold=bold)[0][0]
                width = pdfmetrics.stringWidth(character, font, size)
                if line and line_width + width > content_width:
                    draw_line(line, size=size, bold=bold)
                    line, line_width = "", 0.0
                line += character
                line_width += width
            if line:
                draw_line(line, size=size, bold=bold)

    title = _text(spec.get("title")).strip()
    if title:
        draw_text(title, size=16, bold=True)
        y -= 12
    for block in _blocks(spec):
        kind = str(block.get("type") or "paragraph")
        if kind == "heading":
            level = max(1, min(int(block.get("level") or 2), 6))
            draw_text(block.get("text"), size=14 if level == 1 else 12, bold=True)
            y -= 5
        elif kind == "bullets":
            for item in list(block.get("items") or ())[:10_000]:
                draw_text("- " + _text(item))
            y -= 5
        elif kind == "table":
            for row in list(block.get("rows") or ())[:_MAX_ROWS]:
                if not isinstance(row, (list, tuple)):
                    raise ValueError("table rows must be arrays")
                draw_text(" | ".join(_text(item) for item in row))
            y -= 5
        elif kind == "image":
            payload, _mime = _image_bytes(block)
            with Image.open(BytesIO(payload)) as image:
                pixels_w, pixels_h = image.size
            width = min(_image_width_in(block, 6) * 72, content_width)
            height = width * pixels_h / pixels_w
            max_height = page_height - 2 * margin - 24
            if height > max_height:
                width *= max_height / height
                height = max_height
            if y - height < margin:
                new_page()
            pdf.drawImage(
                ImageReader(BytesIO(payload)), margin, y - height,
                width=width, height=height, mask="auto",
            )
            y -= height + 10
            if block.get("alt"):
                draw_text(block.get("alt"), size=9)
            y -= 5
        elif kind == "page_break":
            new_page()
        elif kind == "quote":
            draw_text(block.get("text"))
            y -= 5
        else:
            draw_text(block.get("text"))
            y -= 5
    pdf.save()
    return buffer.getvalue()


def _tex_escape(value: Any) -> str:
    replacements = {
        "\\": r"\textbackslash{}",
        "{": r"\{",
        "}": r"\}",
        "$": r"\$",
        "&": r"\&",
        "#": r"\#",
        "_": r"\_",
        "%": r"\%",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
    }
    return "".join(replacements.get(character, character) for character in _text(value))


def _latex(spec: Mapping[str, Any]) -> bytes:
    """Build standalone, deterministic LaTeX source from the shared spec.

    Compilation is intentionally a separate renderer/host concern.  Shipping
    valid source is more honest than labeling Markdown bytes as TeX, and lets a
    plugin-provided Tectonic/TeX Live compiler be attached without changing the
    revision identity.
    """

    lines = [
        r"\documentclass[11pt]{article}",
        r"\usepackage[margin=1in]{geometry}",
        r"\usepackage[T1]{fontenc}",
        r"\usepackage{hyperref}",
        r"\usepackage{booktabs}",
        r"\usepackage{longtable}",
        r"\begin{document}",
    ]
    title = _text(spec.get("title")).strip()
    if title:
        lines.extend((r"\title{" + _tex_escape(title) + "}", r"\maketitle"))
    headings = {
        1: "section", 2: "subsection", 3: "subsubsection",
        4: "paragraph", 5: "subparagraph", 6: "subparagraph",
    }
    for block in _blocks(spec):
        kind = str(block.get("type") or "paragraph")
        if kind == "heading":
            command = headings[max(1, min(int(block.get("level") or 2), 6))]
            lines.append(f"\\{command}{{{_tex_escape(block.get('text'))}}}")
        elif kind == "bullets":
            lines.append(r"\begin{itemize}")
            lines.extend(
                r"\item " + _tex_escape(item)
                for item in list(block.get("items") or ())[:10_000]
            )
            lines.append(r"\end{itemize}")
        elif kind == "table":
            rows = [
                list(row) for row in list(block.get("rows") or ())[:_MAX_ROWS]
                if isinstance(row, (list, tuple))
            ]
            if rows:
                width = max(len(row) for row in rows)
                lines.append(r"\begin{longtable}{" + "l" * width + "}")
                for index, row in enumerate(rows):
                    padded = row + [""] * (width - len(row))
                    lines.append(" & ".join(_tex_escape(item) for item in padded) + r" \\")
                    if index == 0:
                        lines.append(r"\midrule")
                lines.append(r"\end{longtable}")
        elif kind == "code":
            lines.extend((r"\begin{verbatim}", _text(block.get("text")), r"\end{verbatim}"))
        elif kind == "quote":
            lines.extend((r"\begin{quote}", _tex_escape(block.get("text")), r"\end{quote}"))
        elif kind == "page_break":
            lines.append(r"\newpage")
        else:
            lines.extend((_tex_escape(block.get("text")), ""))
    lines.extend((r"\end{document}", ""))
    return "\n".join(lines).encode("utf-8")


def _preview_pages(spec: Mapping[str, Any], format_name: str) -> tuple[bytes, ...]:
    from PIL import Image, ImageDraw, ImageFont

    if format_name == "pptx" and isinstance(spec.get("slides"), list):
        pages = []
        for slide in spec["slides"][:100]:
            raw = dict(slide) if isinstance(slide, Mapping) else {}
            text = _text(raw.get("title")) + "\n\n" + "\n".join(
                ("[Image: " + _text(block.get("alt")) + "]")
                if block.get("type") == "image" else _text(block.get("text"))
                for block in raw.get("blocks") or ()
                if isinstance(block, Mapping)
            )
            pages.append(text)
    elif format_name == "xlsx" and isinstance(spec.get("sheets"), list):
        pages = []
        for sheet in spec["sheets"][:50]:
            raw = dict(sheet) if isinstance(sheet, Mapping) else {}
            lines = [_text(raw.get("name") or "Sheet")]
            lines.extend(" | ".join(_text(item) for item in row)
                         for row in list(raw.get("rows") or ())[:35]
                         if isinstance(row, (list, tuple)))
            pages.append("\n".join(lines))
    else:
        # A preview is a concise spec sketch. Never draw a megabyte-long data
        # URI as if it were document text or imply that image pixels vanished.
        preview_spec = dict(spec)
        preview_spec["blocks"] = [
            {"type": "paragraph", "text": "[Image: " + _text(block.get("alt")) + "]"}
            if block.get("type") == "image" else block
            for block in _blocks(spec)
        ]
        lines = _markdown(preview_spec).splitlines()
        pages = ["\n".join(lines[index:index + 45])
                 for index in range(0, max(1, len(lines)), 45)] or [""]
    outputs: list[bytes] = []
    font = ImageFont.load_default()
    for text_value in pages[:100]:
        image = Image.new("RGB", (1200, 800), "white")
        draw = ImageDraw.Draw(image)
        draw.multiline_text((48, 40), text_value[:12_000], fill="#17202a", font=font, spacing=7)
        stream = BytesIO()
        image.save(stream, format="PNG", optimize=True)
        outputs.append(stream.getvalue())
    return tuple(outputs)


def normalized_sha256(format_name: str, payload: bytes) -> str:
    if format_name in {"docx", "pptx", "xlsx"}:
        digest = hashlib.sha256()
        with zipfile.ZipFile(BytesIO(payload), "r") as archive:
            for name in sorted(archive.namelist()):
                digest.update(name.encode("utf-8"))
                digest.update(b"\0")
                digest.update(archive.read(name))
                digest.update(b"\0")
        return digest.hexdigest()
    return hashlib.sha256(payload).hexdigest()


def build_artifact(spec: Mapping[str, Any], format_name: str) -> BuildOutput:
    clean = _strict(dict(spec))
    name = str(format_name or "").strip().lower().lstrip(".")
    builders = {
        "md": (lambda: _markdown(clean).encode("utf-8"), "text/markdown; charset=utf-8", "md", "markdown"),
        "markdown": (lambda: _markdown(clean).encode("utf-8"), "text/markdown; charset=utf-8", "md", "markdown"),
        "html": (lambda: _html(clean).encode("utf-8"), "text/html; charset=utf-8", "html", "html"),
        "json": (lambda: json.dumps(clean, ensure_ascii=False, sort_keys=True, indent=2).encode("utf-8"), "application/json", "json", "json"),
        "csv": (lambda: _delimited(clean, ","), "text/csv; charset=utf-8", "csv", "csv"),
        "tsv": (lambda: _delimited(clean, "\t"), "text/tab-separated-values; charset=utf-8", "tsv", "tsv"),
        "docx": (lambda: _docx(clean), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx", "python-docx"),
        "pptx": (lambda: _pptx(clean), "application/vnd.openxmlformats-officedocument.presentationml.presentation", "pptx", "python-pptx"),
        "xlsx": (lambda: _xlsx(clean), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx", "openpyxl"),
        "pdf": (lambda: _pdf(clean), "application/pdf", "pdf", "reportlab"),
        "latex": (lambda: _latex(clean), "text/x-tex; charset=utf-8", "tex", "latex-source"),
    }
    if name not in builders:
        raise ValueError(f"unsupported artifact format: {format_name}")
    _validate_content(clean, name)
    builder, media_type, extension, renderer = builders[name]
    payload = builder()
    canonical = "md" if name == "markdown" else name
    previews = _preview_pages(clean, canonical)
    return BuildOutput(
        format=canonical,
        payload=payload,
        media_type=media_type,
        extension=extension,
        renderer=renderer,
        renderer_version=BUILDER_VERSION,
        normalized_sha256=normalized_sha256(canonical, payload),
        previews=previews,
        diagnostics=({
            "severity": "info",
            "code": "preview_is_spec_render",
            "message": (
                "PNG previews are deterministic spec sketches, not final PDF or "
                "Microsoft Office fidelity renders."
            ),
        },) if canonical in {"docx", "pptx", "xlsx", "pdf"} else (),
    )


__all__ = ["BUILDER_VERSION", "BuildOutput", "build_artifact", "normalized_sha256"]
