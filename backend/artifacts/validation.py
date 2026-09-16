"""Production structural validation for generated artifact payloads."""

from __future__ import annotations

from collections import Counter
from io import BytesIO, StringIO
import csv
import json
import re
from typing import Any, Mapping


VALIDATOR_VERSION = "variant1-artifact-structural.2"


def _finding(severity: str, code: str, message: str, **location: Any) -> dict[str, Any]:
    return {
        "severity": severity,
        "code": code,
        "message": message,
        "location": {key: value for key, value in location.items() if value is not None},
    }


def _expected_pdf_text(specification: Mapping[str, Any]) -> str:
    """Text content the PDF builder promises to render, excluding decoration."""
    values = [str(specification.get("title") or "")]
    blocks = specification.get("blocks") or []
    if not isinstance(blocks, list):
        return "".join(values)
    for block in blocks:
        if not isinstance(block, Mapping):
            continue
        kind = str(block.get("type") or "paragraph")
        if kind == "bullets":
            values.extend(str(item or "") for item in list(block.get("items") or []))
        elif kind == "table":
            values.extend(str(item or "") for row in list(block.get("rows") or [])
                          if isinstance(row, (list, tuple)) for item in row)
        elif kind == "image":
            values.append(str(block.get("alt") or ""))
        elif kind != "page_break":
            values.append(str(block.get("text") or ""))
    return "\n".join(values)


def validate_payload(
    format_name: str,
    payload: bytes,
    *,
    specification: Mapping[str, Any] | None = None,
) -> tuple[str, tuple[dict[str, Any], ...], dict[str, Any]]:
    name = str(format_name or "").lower().lstrip(".")
    findings: list[dict[str, Any]] = []
    metrics: dict[str, Any] = {"format": name, "bytes": len(payload)}
    if not payload:
        findings.append(_finding("error", "empty_payload", "Artifact payload is empty."))
    try:
        if name == "docx":
            from docx import Document

            document = Document(BytesIO(payload))
            metrics.update({
                "paragraphs": len(document.paragraphs),
                "tables": len(document.tables),
                "sections": len(document.sections),
            })
            if not document.paragraphs and not document.tables:
                findings.append(_finding("error", "docx_no_content", "DOCX has no paragraphs or tables."))
            for index, table in enumerate(document.tables):
                if not table.rows or not table.columns:
                    findings.append(_finding("error", "docx_empty_table", "DOCX contains an empty table.", table=index))
        elif name == "pptx":
            from pptx import Presentation

            presentation = Presentation(BytesIO(payload))
            metrics["slides"] = len(presentation.slides)
            if not presentation.slides:
                findings.append(_finding("error", "pptx_no_slides", "PPTX contains no slides."))
            for slide_index, slide in enumerate(presentation.slides, 1):
                if not slide.shapes:
                    findings.append(_finding("warning", "pptx_blank_slide", "Slide has no shapes.", slide=slide_index))
                for shape_index, shape in enumerate(slide.shapes, 1):
                    if shape.left < 0 or shape.top < 0:
                        findings.append(_finding(
                            "error", "pptx_negative_bounds", "Shape starts outside slide bounds.",
                            slide=slide_index, shape=shape_index,
                        ))
                    if shape.left + shape.width > presentation.slide_width or shape.top + shape.height > presentation.slide_height:
                        findings.append(_finding(
                            "error", "pptx_out_of_bounds", "Shape extends beyond slide bounds.",
                            slide=slide_index, shape=shape_index,
                        ))
        elif name == "xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(BytesIO(payload), data_only=False, read_only=False)
            metrics["sheets"] = len(workbook.worksheets)
            metrics["formulas"] = 0
            if not workbook.worksheets:
                findings.append(_finding("error", "xlsx_no_sheets", "Workbook contains no worksheets."))
            for sheet in workbook.worksheets:
                if sheet.max_row <= 1 and sheet.max_column <= 1 and sheet["A1"].value is None:
                    findings.append(_finding("warning", "xlsx_blank_sheet", "Worksheet is blank.", sheet=sheet.title))
                for row in sheet.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and cell.value.startswith("="):
                            metrics["formulas"] += 1
                            if "#REF!" in cell.value:
                                findings.append(_finding(
                                    "error", "xlsx_broken_reference", "Formula contains #REF!.",
                                    sheet=sheet.title, cell=cell.coordinate,
                                ))
        elif name == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(payload))
            metrics["pages"] = len(reader.pages)
            if not reader.pages:
                findings.append(_finding("error", "pdf_no_pages", "PDF contains no pages."))
            blank = 0
            extracted: list[str] = []
            for index, page in enumerate(reader.pages, 1):
                text = str(page.extract_text() or "").strip()
                extracted.append(text)
                if not text:
                    blank += 1
                    findings.append(_finding("warning", "pdf_blank_page", "PDF page has no extractable text.", page=index))
            metrics["blank_pages"] = blank
            if specification is not None:
                expected = Counter(character for character in _expected_pdf_text(specification)
                                   if not character.isspace())
                actual = Counter(character for character in "\n".join(extracted)
                                 if not character.isspace())
                missing = [(character, count - actual[character])
                           for character, count in expected.items()
                           if actual[character] < count]
                metrics["expected_text_characters"] = sum(expected.values())
                metrics["missing_text_characters"] = sum(count for _character, count in missing)
                for character, count in missing[:20]:
                    findings.append(_finding(
                        "error", "pdf_text_loss",
                        f"PDF lost {count} occurrence(s) of U+{ord(character):04X} from its saved specification.",
                        codepoint=f"U+{ord(character):04X}",
                    ))
                if len(missing) > 20:
                    findings.append(_finding(
                        "error", "pdf_text_loss_more",
                        f"PDF lost text from {len(missing) - 20} additional codepoints.",
                    ))
        elif name in {"json"}:
            decoded = json.loads(payload.decode("utf-8"))
            metrics["root_type"] = type(decoded).__name__
        elif name in {"csv", "tsv"}:
            text = payload.decode("utf-8-sig")
            rows = list(csv.reader(StringIO(text), delimiter="\t" if name == "tsv" else ","))
            metrics["rows"] = len(rows)
            metrics["columns"] = max((len(row) for row in rows), default=0)
            if not rows:
                findings.append(_finding("warning", "tabular_no_rows", "Delimited artifact contains no rows."))
        elif name in {"html"}:
            from lxml import html as lxml_html

            document = lxml_html.fromstring(payload.decode("utf-8"))
            metrics["elements"] = len(document.xpath("//*"))
            if not str(document.text_content() or "").strip():
                findings.append(_finding("warning", "html_no_text", "HTML has no visible text."))
        else:
            text = payload.decode("utf-8")
            metrics["characters"] = len(text)
            if not text.strip():
                findings.append(_finding("warning", "text_no_content", "Text artifact is blank."))
    except Exception as exc:
        findings.append(_finding(
            "error",
            "artifact_parse_failed",
            f"{name or 'artifact'} could not be opened: {type(exc).__name__}: {exc}",
        ))

    if specification is not None:
        serialized = json.dumps(specification, ensure_ascii=False, sort_keys=True)
        dangling = sorted(set(re.findall(r"\{\{cite:([^}]+)\}\}", serialized)))
        for citation in dangling:
            findings.append(_finding(
                "error", "unresolved_citation_placeholder",
                f"Citation placeholder was not compiled: {citation}", citation=citation,
            ))
        metrics["specification_present"] = True
    status = (
        "failed" if any(item["severity"] == "error" for item in findings)
        else "warning" if findings else "passed"
    )
    return status, tuple(findings), metrics


__all__ = ["VALIDATOR_VERSION", "validate_payload"]
